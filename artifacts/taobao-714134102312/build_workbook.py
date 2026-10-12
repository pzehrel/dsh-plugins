import csv
import json
import re
from collections import defaultdict
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path('/Users/pzehrel/Documents/repositories/my/dsh-plugins/artifacts/taobao-714134102312')
SOURCE = Path('/var/folders/4x/dsx65jpx6qq4dkk4h7hz9j8m0000gn/T/dsh-spill-YFUDct/session-2b4cc1e5aea3/12ebff8526d1-browser_execute.txt')
ROOT.mkdir(parents=True, exist_ok=True)
text = SOURCE.read_text(encoding='utf-8')
decoder = json.JSONDecoder()
metadata, _ = decoder.raw_decode(text.split('"metadata":', 1)[1])
raw_rows, _ = decoder.raw_decode(text.split('"rows":', 1)[1])
assert len(raw_rows) == metadata['skuCount'] == metadata['propertyCount'] == 250
assert len({r[2] for r in raw_rows}) == 250
headers = ['页面序号', '类别', 'SKU名称', '螺纹规格', '公称直径（mm）', '长度（mm）', '每包数量（个）', '优惠前（元/包）', '店铺优惠后（元/包）', '优惠前折算（元/个）', '优惠后折算（元/个）', '省下（元/包）', '库存状态', 'SKU ID', '规格属性路径', '商品链接', '备注']
rows = []
groups = defaultdict(list)
for order, name, sku_id, prop_path, before_fen, before_text, after_fen, after_text, status in raw_rows:
    before = Decimal(before_fen) / 100
    after = Decimal(after_fen) / 100
    assert before == Decimal(before_text)
    assert after == Decimal(after_text)
    assert Decimal(0) < after <= before
    match = re.fullmatch(r'M([\d.]+)\*(\d+)\((\d+)个\)', name)
    if match:
        diameter, length, count = Decimal(match[1]), int(match[2]), int(match[3])
        thread = 'M' + match[1]
        category = '螺丝'
        unit_before = float((before / count).quantize(Decimal('0.000001')))
        unit_after = float((after / count).quantize(Decimal('0.000001')))
        note = ''
    else:
        diameter = length = count = unit_before = unit_after = None
        thread = ''
        category = '套装' if '套装' in name else '其他选项'
        note = '名称未标明总颗数，不计算每颗单价。' if category == '套装' else '宣传文案选项，页面有独立SKU和价格；实际内容未注明。'
    url = f'https://item.taobao.com/item.htm?id={metadata["itemId"]}&skuId={sku_id}'
    row = [order, category, name, thread, float(diameter) if diameter is not None else None, length, count, float(before), float(after), unit_before, unit_after, float(before - after), status, sku_id, prop_path, url, note]
    rows.append(row)
    groups[thread if match else category].append(row)
assert sum(r[1] == '螺丝' for r in rows) == 246
assert sum(r[1] == '套装' for r in rows) == 3
assert sum(r[1] == '其他选项' for r in rows) == 1

collected = datetime.fromisoformat(metadata['collectedAt']).astimezone(ZoneInfo('Asia/Shanghai'))
summary_headers = ['规格/类别', 'SKU数', '最短长度（mm）', '最长长度（mm）', '每包数量', '优惠前最低（元/包）', '优惠前最高（元/包）', '优惠后最低（元/包）', '优惠后最高（元/包）', '优惠后最低（元/个）', '优惠后最高（元/个）']
summary = []
for group, values in groups.items():
    lengths = [r[5] for r in values if r[5] is not None]
    counts = sorted({r[6] for r in values if r[6] is not None})
    units = [r[10] for r in values if r[10] is not None]
    summary.append([group, len(values), min(lengths) if lengths else None, max(lengths) if lengths else None, '、'.join(map(str, counts)), min(r[7] for r in values), max(r[7] for r in values), min(r[8] for r in values), max(r[8] for r in values), min(units) if units else None, max(units) if units else None])
summary.sort(key=lambda r: (0, float(r[0][1:])) if r[0].startswith('M') else (1, r[0]))

wb = Workbook()
ws = wb.active
ws.title = '全部SKU'
ws.append(headers)
for row in rows:
    ws.append(row)
ws.freeze_panes = 'D2'
table = Table(displayName='AllSKU', ref=f'A1:Q{len(rows)+1}')
table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
ws.add_table(table)
for cells in ws.iter_rows(min_row=2):
    for col in (8, 9, 12):
        cells[col-1].number_format = '0.00'
    for col in (10, 11):
        cells[col-1].number_format = '0.000000'
    cells[13].number_format = '@'
    cells[15].hyperlink = cells[15].value
    cells[15].style = 'Hyperlink'
widths = [12, 12, 44, 12, 19, 14, 17, 20, 24, 25, 25, 19, 13, 20, 29, 58, 65]
for i, width in enumerate(widths, 1):
    ws.column_dimensions[ws.cell(1, i).column_letter].width = width

ss = wb.create_sheet('规格汇总')
ss.append(summary_headers)
for row in summary:
    ss.append(row)
ss.freeze_panes = 'B2'
ss.auto_filter.ref = ss.dimensions
for i in range(1, 12):
    ss.column_dimensions[ss.cell(1, i).column_letter].width = 23 if i >= 6 else 19
for cells in ss.iter_rows(min_row=2):
    for col in range(6, 10):
        cells[col-1].number_format = '0.00'
    for col in (10, 11):
        cells[col-1].number_format = '0.000000'

special = wb.create_sheet('套装及其他')
special.append(headers)
for row in rows:
    if row[1] != '螺丝':
        special.append(row)
special.freeze_panes = 'D2'
special.auto_filter.ref = special.dimensions
for i, width in enumerate(widths, 1):
    special.column_dimensions[special.cell(1, i).column_letter].width = width
for cells in special.iter_rows(min_row=2):
    for col in (8, 9, 12):
        cells[col-1].number_format = '0.00'
    cells[13].number_format = '@'
    cells[15].hyperlink = cells[15].value
    cells[15].style = 'Hyperlink'

notes = wb.create_sheet('采集说明')
notes.append(['项目', '内容'])
note_rows = [
    ['商品', metadata['title']],
    ['店铺', metadata['shop']],
    ['商品ID', metadata['itemId']],
    ['来源', f'https://item.taobao.com/item.htm?id={metadata["itemId"]}'],
    ['采集时间（北京时间）', collected.strftime('%Y-%m-%d %H:%M:%S')],
    ['SKU总数', 250],
    ['类别统计', '246个螺丝规格 + 3个套装 + 1个宣传文案选项；全部保留。'],
    ['提取来源', '登录后的商品页：skuBase.props.values、skuBase.skus与skuCore.sku2info；按规格属性路径匹配SKU ID。'],
    ['优惠前', '页面标注“优惠前”的每包售价；不将其称为历史原价或划线原价。'],
    ['店铺优惠后', '当前登录账号看到的“店铺优惠后”每包售价；页面显示官方立减12%和券满2减1。'],
    ['优惠使用', '页面按钮为“领券购买”。优惠是否实际可用、是否已领取及最终应付价，以领取和下单页为准；未替用户领券或下单。'],
    ['包装数量', '从SKU名称中的“(N个)”提取，表示该售价对应的一包数量；不等同于库存。'],
    ['折算单价', '每包售价除以每包数量，保留6位小数；便于比较，不是可单颗购买的售价。'],
    ['套装数量', '10格、15格、24格是盒格数；商品页未在SKU名称中注明总颗数，未据此推算数量或单价。'],
    ['库存', '所有SKU均显示“有货”。不输出接口中的quantity作为实际库存：部分记录同时标有moreQuantity，数量可能为展示或购买上限。'],
    ['完整性核对', '250个选项对应250个唯一SKU ID；无价格缺失；价格分值与显示文本逐项一致。'],
    ['页面抽查', 'M4*50(20个)：3.85 / 2.39；M1.6*8(100个)：2.20 / 0.94；24格大盒套装：34.10 / 29.00；M16*100(1个)：7.71 / 5.78。顺序为优惠前 / 店铺优惠后，单位元/包。'],
    ['时效', '价格是采集时的页面快照，可能随促销、账号权益和下单条件变化；运费未计入单价。'],
]
for row in note_rows:
    notes.append(row)
notes.column_dimensions['A'].width = 27
notes.column_dimensions['B'].width = 110
notes.freeze_panes = 'B2'
for cells in notes.iter_rows(min_row=2):
    cells[1].alignment = Alignment(wrap_text=True, vertical='top')
    notes.row_dimensions[cells[0].row].height = 42 if len(str(cells[1].value)) > 70 else 28
notes['B4'].number_format = '@'
notes['B5'].hyperlink = notes['B5'].value
notes['B5'].style = 'Hyperlink'
for sheet in wb:
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 30
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor='17365D')
        cell.font = Font(name='微软雅黑', bold=True, color='FFFFFF')
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if not cell.hyperlink:
                cell.font = Font(name='微软雅黑', size=10)
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

xlsx_path = ROOT / '淘宝304沉头内六角螺丝_全部250个SKU及价格.xlsx'
csv_path = ROOT / '淘宝304沉头内六角螺丝_全部250个SKU及价格.csv'
json_path = ROOT / 'sku_source.json'
wb.save(xlsx_path)
with csv_path.open('w', encoding='utf-8-sig', newline='') as file:
    writer = csv.writer(file)
    writer.writerow(headers)
    writer.writerows(rows)
json_path.write_text(json.dumps({'metadata': metadata, 'columns': ['order', 'name', 'skuId', 'propertyPath', 'priceBeforeFen', 'priceBeforeText', 'priceAfterFen', 'priceAfterText', 'quantityText'], 'rows': raw_rows}, ensure_ascii=False, indent=2), encoding='utf-8')
check = load_workbook(xlsx_path, data_only=True)
assert check['全部SKU'].max_row == 251
assert check['套装及其他'].max_row == 5
assert len({check['全部SKU'].cell(i, 14).value for i in range(2, 252)}) == 250
for i, row in enumerate(rows, 2):
    assert check['全部SKU'].cell(i, 3).value == row[2]
    assert check['全部SKU'].cell(i, 8).value == row[7]
    assert check['全部SKU'].cell(i, 9).value == row[8]
    assert check['全部SKU'].cell(i, 14).value == row[13]
assert all(check['全部SKU'].cell(i, 13).value == '有货' for i in range(2, 252))
result = {'xlsx': str(xlsx_path), 'csv': str(csv_path), 'collectedAt': collected.isoformat(), 'rows': len(rows), 'screwRows': 246, 'sets': 3, 'otherRows': 1, 'minAfter': min(r[8] for r in rows), 'maxAfter': max(r[8] for r in rows), 'summary': summary, 'specialRows': [r for r in rows if r[1] != '螺丝']}
(ROOT / 'verification.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(result, ensure_ascii=False, indent=2))
